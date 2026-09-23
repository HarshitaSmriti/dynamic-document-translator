"""
Document language detection based on robust Unicode character script distributions.
Analyzes entire document text to detect English, Hindi, or Bengali.
"""

from dataclasses import dataclass
from typing import Optional, Dict
from docx import Document


@dataclass
class DetectionResult:
    """
    Result of language detection with script counts and confidence metrics.
    """
    detected_language: Optional[str]
    confidence: float
    char_counts: Dict[str, int]
    total_valid_chars: int
    is_ambiguous: bool
    warning_message: Optional[str] = None


def count_script_characters(text: str) -> Dict[str, int]:
    """
    Counts characters belonging to English (Latin), Hindi (Devanagari), and Bengali scripts.
    """
    counts = {
        "English": 0,
        "Hindi": 0,
        "Bengali": 0,
    }

    for char in text:
        code = ord(char)
        # Devanagari script range: U+0900 - U+097F
        if 0x0900 <= code <= 0x097F:
            counts["Hindi"] += 1
        # Bengali script range: U+0980 - U+09FF
        elif 0x0980 <= code <= 0x09FF:
            counts["Bengali"] += 1
        # Latin alphabet (English): a-z, A-Z
        elif ("a" <= char <= "z") or ("A" <= char <= "Z"):
            counts["English"] += 1

    return counts


def detect_language_from_text(text: str, min_chars_threshold: int = 10) -> DetectionResult:
    """
    Detects language from text string by analyzing script character frequencies.
    """
    counts = count_script_characters(text)
    total = sum(counts.values())

    if total < min_chars_threshold:
        return DetectionResult(
            detected_language=None,
            confidence=0.0,
            char_counts=counts,
            total_valid_chars=total,
            is_ambiguous=True,
            warning_message=(
                f"The document contains too little recognizable text ({total} characters). "
                f"A minimum of {min_chars_threshold} alphabetic characters is required for accurate detection."
            ),
        )

    # Sort candidates by count descending
    sorted_candidates = sorted(counts.items(), key=lambda item: item[1], reverse=True)
    top_lang, top_count = sorted_candidates[0]
    second_lang, second_count = sorted_candidates[1]

    confidence = top_count / total

    # Check for mixed or ambiguous language content
    is_ambiguous = False
    warning_message = None

    if confidence < 0.65 and second_count > 0:
        is_ambiguous = True
        warning_message = (
            f"The document contains mixed languages ({top_lang}: {top_count} chars, "
            f"{second_lang}: {second_count} chars). Detected primary language is {top_lang}."
        )

    return DetectionResult(
        detected_language=top_lang,
        confidence=confidence,
        char_counts=counts,
        total_valid_chars=total,
        is_ambiguous=is_ambiguous,
        warning_message=warning_message,
    )


def extract_full_text_from_docx(doc: Document) -> str:
    """
    Extracts all text content from paragraphs, headers, footers, and tables in a DOCX Document.
    """
    text_parts = []

    # Body paragraphs
    for p in doc.paragraphs:
        if p.text:
            text_parts.append(p.text)

    # Tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if p.text:
                        text_parts.append(p.text)

    # Sections (Headers and Footers)
    for section in doc.sections:
        if section.header:
            for p in section.header.paragraphs:
                if p.text:
                    text_parts.append(p.text)
        if section.footer:
            for p in section.footer.paragraphs:
                if p.text:
                    text_parts.append(p.text)

    return "\n".join(text_parts)


def detect_document_language(doc: Document) -> DetectionResult:
    """
    Detects language of a python-docx Document object.
    """
    extracted_text = extract_full_text_from_docx(doc)
    return detect_language_from_text(extracted_text)
